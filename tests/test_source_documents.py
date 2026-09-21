from __future__ import annotations

import struct
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.document import Document as WordDocument
from openpyxl import Workbook
from PIL import Image

import metaculus_bot.research.source_documents as source_documents
from metaculus_bot.research.source_documents import (
    SourceReadError,
    is_local_source,
    parse_source,
)


def _xlsx_bytes(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _docx_bytes(document: WordDocument) -> bytes:
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _zip_bytes(members: dict[str, bytes]) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, contents in members.items():
            archive.writestr(name, contents)
    return output.getvalue()


def _mark_first_zip_member_encrypted(body: bytes) -> bytes:
    """Set the ZIP encryption flag without encrypting data; detection must fail shut."""
    marked = bytearray(body)
    local_header = marked.index(b"PK\x03\x04")
    local_flags = struct.unpack_from("<H", marked, local_header + 6)[0]
    struct.pack_into("<H", marked, local_header + 6, local_flags | 1)
    central_header = marked.index(b"PK\x01\x02")
    central_flags = struct.unpack_from("<H", marked, central_header + 8)[0]
    struct.pack_into("<H", marked, central_header + 8, central_flags | 1)
    return bytes(marked)


def test_is_local_source_detects_supported_signatures_and_mimes() -> None:
    workbook = Workbook()
    xlsx = _xlsx_bytes(workbook)
    document = Document()
    docx = _docx_bytes(document)
    archive = _zip_bytes({"notes.txt": b"plain text"})

    assert is_local_source(xlsx, "application/octet-stream")
    assert is_local_source(docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert is_local_source(archive, "application/zip")
    assert is_local_source(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest", "application/vnd.ms-excel")
    assert not is_local_source(b"PK is the first two letters of this ordinary CSV row", "text/csv")
    assert not is_local_source(b"<html><body>page</body></html>", "text/html")
    assert not is_local_source(b"not a container", "application/octet-stream")


def test_archive_reads_csv_tsv_literally_and_keeps_unknown_members() -> None:
    body = _zip_bytes(
        {
            "data/results.csv": b"region,value\nNorth,<b>literal</b>\nSouth,late-row\n",
            "data/notes.tsv": b"field\tvalue\nraw\t<x>literal</x>\n",
            "payload.bin": b"\x00\x01\x02",
            "nested.zip": _zip_bytes({"ignored.txt": b"must not recurse"}),
        }
    )

    parsed = parse_source(body, "application/zip")

    assert parsed.kind == "archive"
    assert len(parsed.sections) == 2
    csv_section = next(section for section in parsed.sections if section.member == "data/results.csv")
    tsv_section = next(section for section in parsed.sections if section.member == "data/notes.tsv")
    assert "row 3" in csv_section.text
    assert "<b>literal</b>" in csv_section.text
    assert "<x>literal</x>" in tsv_section.text
    assert "ignored.txt" not in "\n".join(section.text for section in parsed.sections)

    member_map = {member.name: member for member in parsed.members}
    assert member_map["data/results.csv"].readable
    assert member_map["data/notes.tsv"].readable
    assert not member_map["payload.bin"].readable
    assert member_map["payload.bin"].reason
    assert not member_map["nested.zip"].readable
    assert member_map["nested.zip"].reason
    assert parsed.retained_bytes >= sum(len(section.text.encode("utf-8")) for section in parsed.sections)


def test_archive_keeps_directory_inventory_but_marks_it_unreadable() -> None:
    parsed = parse_source(
        _zip_bytes({"reports/": b"", "reports/summary.txt": b"Total revenue: 125"}),
        "application/zip",
    )

    directory = next(member for member in parsed.members if member.name == "reports/")
    assert not directory.readable
    assert directory.reason == "directory has no file content"
    assert [section.member for section in parsed.sections] == ["reports/summary.txt"]


def test_csv_sniffs_semicolon_delimiter_without_splitting_decimal_or_quoted_cells() -> None:
    body = _zip_bytes(
        {
            "Europe.csv": (
                b'country;share;note\r\nCH;14,09;"Reference says; the rate is 14,09 percent, approximately."\r\n'
            )
        }
    )

    parsed = parse_source(body, "application/zip")

    assert parsed.sections[0].text.splitlines() == [
        "columns: A B C",
        "row 1 (A1:C1): country\tshare\tnote",
        "row 2 (A2:C2): CH\t14,09\tReference says; the rate is 14,09 percent, approximately.",
    ]


def test_csv_sniff_ignores_but_keeps_leading_metadata_comments() -> None:
    columns = ["drought_region_id", "measured_at"] + [f"metric_{index}" for index in range(18)]
    values = ["31", "14.09.2026", "1,25", '"south; central, 1,25"'] + ["0"] * 16
    csv_text = "\n".join(
        [
            "# last updated: 20.09.2026",
            "# https://example.test/drought.csv",
            "# public data",
            ";".join(columns),
            ";".join(values),
        ]
    )
    body = _zip_bytes({"weekly_current_regions.csv": csv_text.encode()})

    parsed = parse_source(body, "application/zip")

    section = parsed.sections[0]
    assert section.columns == 20
    assert "row 1 (A1:A1): # last updated: 20.09.2026" in section.text
    assert "row 4 (A4:T4): drought_region_id" in section.text
    assert "row 5 (A5:T5): 31\t14.09.2026\t1,25\tsouth; central, 1,25" in section.text


def test_archive_parses_office_member_with_member_identity() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Summary"
    sheet["A1"] = "Quarter"
    sheet["B1"] = "Q3"
    sheet["A2"] = "Revenue"
    sheet["B2"] = 125
    body = _zip_bytes({"reports/quarter.xlsx": _xlsx_bytes(workbook), "readme.bin": b"Report archive"})

    parsed = parse_source(body, "application/zip")

    assert parsed.kind == "archive"
    assert len(parsed.sections) == 1
    assert parsed.sections[0].member == "reports/quarter.xlsx"
    assert parsed.sections[0].sheet == "Summary"
    assert "A2=Revenue" in parsed.sections[0].text
    assert "B2=125" in parsed.sections[0].text
    assert all(member.readable for member in parsed.members if member.name == "reports/quarter.xlsx")


def test_xlsx_reads_all_sheets_late_rows_formats_and_formula_cache_status() -> None:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Forecast"
    first["A1"] = "Date"
    first["A2"] = 46_000
    first["A2"].number_format = "yyyy-mm-dd"
    first["B1"] = "Growth"
    first["B2"] = 0.125
    first["B2"].number_format = "0.0%"
    first["C1"] = "Calculated"
    first["C2"] = "=B2*2"
    first.merge_cells("D1:E1")
    first["D1"] = "Merged heading"
    first["A1200"] = "late row value"
    workbook.create_sheet("Second sheet")["B3"] = "second-sheet value"

    parsed = parse_source(_xlsx_bytes(workbook), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    assert parsed.kind == "workbook"
    assert [section.sheet for section in parsed.sections] == ["Forecast", "Second sheet"]
    forecast = parsed.sections[0].text
    assert "A2=2025-12-09" in forecast
    assert "B2=12.5%" in forecast
    assert "C2==B2*2" in forecast
    assert "cached value unavailable" in forecast
    assert any("formulas are not evaluated" in notice for notice in parsed.notices)
    assert "A1200=late row value" in forecast
    assert "E1=[blank]" in forecast
    assert "second-sheet value" in parsed.sections[1].text
    assert parsed.sections[0].rows == 1200
    assert parsed.sections[0].columns == 5


@pytest.mark.parametrize(
    ("number_format", "formatted_cell"),
    [
        ('0.0"%"', 'A1=0.125 [literal percent format: 0.0"%"]'),
        (r"0.0\%", r"A1=0.125 [literal percent format: 0.0\%]"),
    ],
)
def test_xlsx_literal_percent_format_does_not_rescale_numeric_value(
    number_format: str,
    formatted_cell: str,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = 0.125
    sheet["A1"].number_format = number_format

    parsed = parse_source(_xlsx_bytes(workbook), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    assert formatted_cell in parsed.sections[0].text


def test_xlsx_enforces_cell_limit_without_truncating(monkeypatch: pytest.MonkeyPatch) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "one"
    sheet["B1"] = "two"
    monkeypatch.setattr(source_documents, "LOCAL_SOURCE_MAX_CELLS", 1)
    with pytest.raises(SourceReadError, match="cell"):
        parse_source(_xlsx_bytes(workbook), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def test_xlsx_enforces_character_limit_without_truncating(monkeypatch: pytest.MonkeyPatch) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "one"
    sheet["B1"] = "two"
    monkeypatch.setattr(source_documents, "LOCAL_SOURCE_MAX_CHARS", 5)
    with pytest.raises(SourceReadError, match="character"):
        parse_source(_xlsx_bytes(workbook), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def test_xlsx_enforces_sheet_limit_without_truncating(monkeypatch: pytest.MonkeyPatch) -> None:
    workbook = Workbook()
    workbook.create_sheet("Second sheet")
    monkeypatch.setattr(source_documents, "LOCAL_SOURCE_MAX_SHEETS", 1)

    with pytest.raises(SourceReadError, match="sheet"):
        parse_source(_xlsx_bytes(workbook), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def test_word_preserves_paragraph_table_order_nested_tables_and_headers() -> None:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Header label"
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Outer cell"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested cell"
    document.add_paragraph("After table")
    document.sections[0].footer.paragraphs[0].text = "Footer label"

    parsed = parse_source(
        _docx_bytes(document),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert parsed.kind == "word"
    section_text = "\n".join(section.text for section in parsed.sections)
    assert section_text.index("Before table") < section_text.index("Outer cell")
    assert section_text.index("Outer cell") < section_text.index("Nested cell")
    assert section_text.index("Nested cell") < section_text.index("After table")
    assert "Header label" in section_text
    assert "Footer label" in section_text
    assert any("header" in (section.sheet or "").lower() for section in parsed.sections)
    assert any("footer" in (section.sheet or "").lower() for section in parsed.sections)


def test_word_reads_default_first_page_and_even_page_story_variants() -> None:
    document = Document()
    document.settings.odd_and_even_pages_header_footer = True
    section = document.sections[0]
    section.different_first_page_header_footer = True
    stories = (
        ("header section 1", section.header, "Default header"),
        ("first-page header section 1", section.first_page_header, "First-page header"),
        ("even-page header section 1", section.even_page_header, "Even-page header"),
        ("footer section 1", section.footer, "Default footer"),
        ("first-page footer section 1", section.first_page_footer, "First-page footer"),
        ("even-page footer section 1", section.even_page_footer, "Even-page footer"),
    )
    for _label, story, text in stories:
        story.paragraphs[0].text = text

    parsed = parse_source(
        _docx_bytes(document),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    story_sections = [section for section in parsed.sections if section.sheet != "body"]
    assert [section.sheet for section in story_sections] == [label for label, _story, _text in stories]
    assert [section.text for section in story_sections] == [
        f"[{label}]\nparagraph 1: {text}" for label, _story, text in stories
    ]


def test_word_discloses_omitted_visual_content() -> None:
    image = Image.new("RGB", (1, 1), color="white")
    image_bytes = BytesIO()
    image.save(image_bytes, format="PNG")
    document = Document()
    paragraph = document.add_paragraph("Figure: ")
    paragraph.add_run().add_picture(BytesIO(image_bytes.getvalue()))

    parsed = parse_source(
        _docx_bytes(document),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert any("visual content omitted" in notice.lower() for notice in parsed.notices)
    assert any("visual content omitted" in section.text.lower() for section in parsed.sections)


def test_archive_limits_use_actual_expansion_and_entry_count(monkeypatch: pytest.MonkeyPatch) -> None:
    body = _zip_bytes({"small.csv": b"x" * 64})
    monkeypatch.setattr(source_documents, "LOCAL_SOURCE_MAX_EXPANDED_BYTES", 32)
    with pytest.raises(SourceReadError, match="expanded"):
        parse_source(body, "application/zip")

    monkeypatch.setattr(source_documents, "LOCAL_SOURCE_MAX_EXPANDED_BYTES", 20 * 1024 * 1024)
    monkeypatch.setattr(source_documents, "LOCAL_SOURCE_MAX_ENTRIES", 1)
    many_entries = _zip_bytes({"one.txt": b"1", "two.txt": b"2"})
    with pytest.raises(SourceReadError, match="entry"):
        parse_source(many_entries, "application/zip")


@pytest.mark.parametrize(
    ("body", "content_type", "reason_fragment"),
    [
        (b"PK\x03\x04broken", "application/zip", "archive"),
        (_mark_first_zip_member_encrypted(_zip_bytes({"private.csv": b"secret"})), "application/zip", "encrypt"),
        (b"broken xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "workbook"),
        (b"not a document", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "word"),
    ],
)
def test_malformed_or_encrypted_source_fails_with_specific_reason(
    body: bytes,
    content_type: str,
    reason_fragment: str,
) -> None:
    with pytest.raises(SourceReadError) as raised:
        parse_source(body, content_type)

    assert raised.value.reason
    assert reason_fragment in raised.value.reason.lower()


def test_unrecognized_source_is_refused() -> None:
    with pytest.raises(SourceReadError, match="unsupported"):
        parse_source(b"ordinary body", "text/plain")


def test_xls_fixture_is_read_with_cached_values_and_number_formats() -> None:
    fixture = Path(__file__).parent / "fixtures" / "source_documents" / "legacy.xls"

    parsed = parse_source(fixture.read_bytes(), "application/vnd.ms-excel")

    assert parsed.kind == "workbook"
    assert any(section.sheet == "Sales 2025" for section in parsed.sections)
    text = "\n".join(section.text for section in parsed.sections)
    assert "2025-01-31" in text
    assert "12.5%" in text
    assert "saved workbook values" in text.lower()


def test_xls_signature_wins_over_misleading_mime_type() -> None:
    fixture = Path(__file__).parent / "fixtures" / "source_documents" / "legacy.xls"

    parsed = parse_source(fixture.read_bytes(), "text/html")

    assert parsed.kind == "workbook"
    assert "2025-01-31" in parsed.sections[0].text
