"""Offline acceptance coverage for local source formats and same-agent image evidence."""

from __future__ import annotations

import copy
import json
import re
from collections import Counter
from collections.abc import AsyncIterator, Iterator
from contextlib import asynccontextmanager
from io import BytesIO
from typing import Any
from unittest.mock import MagicMock
from zipfile import ZIP_DEFLATED, ZipFile

import aiohttp
import httpx
import litellm
import pytest
from aiohttp import web
from aiohttp.test_utils import TestServer
from docx import Document
from docx.document import Document as DocumentType
from litellm.llms.custom_httpx.http_handler import AsyncHTTPHandler
from openpyxl import Workbook
from PIL import Image

from metaculus_bot.constants import GAP_FILL_V2_DRIVER_MODEL
from metaculus_bot.research.agentic import llm as agentic_llm
from metaculus_bot.research.agentic import tools as agentic_tools
from metaculus_bot.research.agentic.image_messages import bind_image_materialization, image_reference_message
from metaculus_bot.research.agentic.loop import run_agentic_loop
from metaculus_bot.research.agentic.types import ToolSpec
from metaculus_bot.research.fetch_ladder import guard, run_cache
from metaculus_bot.research.fetch_ladder.context import LadderContext
from metaculus_bot.research.fetch_ladder.ladder import fetch_url
from metaculus_bot.research.fetch_ladder.policy import RESOLUTION_SOURCE_POLICY
from metaculus_bot.research.image_assets import ImageView
from metaculus_bot.research.source_presentation import select_source_sections, source_text
from tests.agentic_fakes import FakeLlm
from tests.agentic_fakes import gap_accounting as _accounting
from tests.agentic_fakes import loop_config as _config
from tests.agentic_fakes import plan_call as _plan_call
from tests.agentic_fakes import response as _response
from tests.agentic_fakes import tool_call as _tool_call

_CONTINUATION_RE = re.compile(
    r"\n\[truncated at (?P<offset>\d+) of (?P<total>\d+) chars — call again with start_char=(?P=offset)\]"
)


@pytest.fixture(autouse=True)
def _isolated_source_cache() -> Iterator[None]:
    run_cache.clear()
    yield
    run_cache.clear()


def _zip_bytes(members: dict[str, bytes]) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, contents in members.items():
            archive.writestr(name, contents)
    return output.getvalue()


def _xlsx_bytes(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _docx_bytes(document: DocumentType) -> bytes:
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _png_bytes() -> bytes:
    image = Image.new("RGB", (12, 8), color="navy")
    for x in range(6, 12):
        for y in range(8):
            image.putpixel((x, y), (255, 220, 0))
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _solid_png_bytes(color: tuple[int, int, int]) -> bytes:
    image = Image.new("RGB", (6, 4), color=color)
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


@asynccontextmanager
async def _local_source_server(
    routes: dict[str, tuple[bytes, str]],
) -> AsyncIterator[tuple[dict[str, str], Counter[str]]]:
    hits: Counter[str] = Counter()

    async def serve(request: web.Request) -> web.Response:
        hits[request.path] += 1
        body, content_type = routes[request.path]
        return web.Response(body=body, headers={"Content-Type": content_type})

    app = web.Application()
    app.router.add_get("/{path:.*}", serve)
    server = TestServer(app)
    await server.start_server()
    try:
        urls = {path: str(server.make_url(path)) for path in routes}
        yield urls, hits
    finally:
        await server.close()


def _tool_content(result: Any, tool_call_id: str) -> str:
    return next(
        message["content"]
        for message in result.transcript
        if message.get("role") == "tool" and message.get("tool_call_id") == tool_call_id
    )


def _tool(name: str, tools: list[ToolSpec]) -> ToolSpec:
    return next(tool for tool in tools if tool.name == name)


@pytest.mark.asyncio
async def test_archive_inventory_then_late_member_query_banks_grounded_finding_without_redownload(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    late_quote = "row 2402 (A2402:C2402): 2026-09-20\t947\tfinal official hospitalizations"
    csv_rows = ["date,count,note", *(f"2025-01-{index % 28 + 1:02d},{index},historical" for index in range(2400))]
    csv_rows.append("2026-09-20,947,final official hospitalizations")
    archive = _zip_bytes(
        {
            "notes.txt": b"Background only.",
            "data/results.csv": ("\n".join(csv_rows) + "\n").encode(),
        }
    )
    paid_reader = MagicMock(side_effect=AssertionError("local formats must never reach the paid reader"))
    monkeypatch.setattr(agentic_tools, "_run_document_read_sync", paid_reader)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    monkeypatch.setattr(guard, "is_public_http_url", lambda _url: _async_true())

    async with (
        _local_source_server(
            {
                "/data.zip": (archive, "application/zip"),
                "/broken.zip": (b"PK\x03\x04broken", "application/zip"),
            }
        ) as (urls, hits),
        aiohttp.ClientSession() as session,
    ):
        tools = agentic_tools.build_gap_fill_tools(
            "September hospitalizations", ctx=agentic_tools.question_ladder_context(session=session)
        )
        fake_llm = FakeLlm(
            [
                _response(tool_calls=[_plan_call()]),
                _response(tool_calls=[_tool_call("inventory", "fetch", {"url": urls["/data.zip"]})]),
                _response(
                    tool_calls=[
                        _tool_call(
                            "query",
                            "read_document",
                            {
                                "url": urls["/data.zip"],
                                "ask": "official hospitalizations",
                                "member": "data/results.csv",
                            },
                        )
                    ]
                ),
                _response(
                    tool_calls=[
                        _tool_call(
                            "record",
                            "record_findings",
                            {
                                "findings": [
                                    {
                                        "claim": "The final official count is 947.",
                                        "source_url": urls["/data.zip"],
                                        "quote": late_quote,
                                        "retrieved_how": "read_document",
                                    }
                                ]
                            },
                        )
                    ]
                ),
                _response(tool_calls=[_tool_call("done", "conclude", {"gap_accounting": _accounting("g1")})]),
            ]
        )

        result = await run_agentic_loop(
            "system",
            "brief",
            tools,
            _config(max_steps=6, max_conclude_gate_rejections=0),
            llm_call=fake_llm,
        )

        refusal = await _tool("read_document", tools).handler(
            url=urls["/data.zip"], ask="anything", member="missing.csv"
        )
        malformed = await _tool("read_document", tools).handler(url=urls["/broken.zip"], ask="anything")
        cached_inventory = await _tool("fetch", tools).handler(url=urls["/data.zip"])

    inventory = _tool_content(result, "inventory")
    query = _tool_content(result, "query")
    assert "method: local_navigation" in inventory
    assert "data/results.csv" in inventory
    assert "final official hospitalizations" not in inventory
    assert "method: digest_local" in query
    assert late_quote in query
    assert result.telemetry.findings_count == 1
    assert "The final official count is 947." in result.findings_markdown
    assert hits["/data.zip"] == 1
    assert hits["/broken.zip"] == 1
    assert cached_inventory.method == "local_navigation"
    assert "final official hospitalizations" not in cached_inventory.content_markdown
    assert refusal.status == "error"
    assert malformed.status == "error"
    assert all(len(_tool_content(result, call_id)) <= 8_000 for call_id in ("inventory", "query", "record"))
    paid_reader.assert_not_called()


@pytest.mark.asyncio
async def test_workbook_sheet_selector_uses_cached_download_and_excludes_other_sheets(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = Workbook()
    overview = workbook.active
    assert overview is not None
    overview.title = "Overview"
    overview["A1"] = "This sheet is only an overview"
    evidence = workbook.create_sheet("Evidence")
    evidence["A1"] = "Region"
    evidence["B1"] = "Hospitalizations"
    evidence["A1200"] = "North"
    evidence["B1200"] = 947
    monkeypatch.setattr(guard, "is_public_http_url", lambda _url: _async_true())

    async with (
        _local_source_server(
            {
                "/report.xlsx": (
                    _xlsx_bytes(workbook),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            }
        ) as (urls, hits),
        aiohttp.ClientSession() as session,
    ):
        tools = agentic_tools.build_gap_fill_tools(
            "North hospitalizations", ctx=agentic_tools.question_ladder_context(session=session)
        )
        fake_llm = FakeLlm(
            [
                _response(tool_calls=[_plan_call()]),
                _response(tool_calls=[_tool_call("inventory", "fetch", {"url": urls["/report.xlsx"]})]),
                _response(
                    tool_calls=[
                        _tool_call(
                            "sheet",
                            "read_document",
                            {
                                "url": urls["/report.xlsx"],
                                "ask": "North Hospitalizations",
                                "sheet": "Evidence",
                            },
                        )
                    ]
                ),
                _response(tool_calls=[_tool_call("done", "conclude", {"gap_accounting": _accounting("g1")})]),
            ]
        )
        result = await run_agentic_loop(
            "system",
            "brief",
            tools,
            _config(max_steps=5, max_conclude_gate_rejections=0),
            llm_call=fake_llm,
        )

    inventory = _tool_content(result, "inventory")
    selected = _tool_content(result, "sheet")
    assert "method: local_navigation" in inventory
    assert "Overview" in inventory
    assert "Evidence" in inventory
    assert "North" not in inventory
    assert "method: digest_local" in selected
    assert "Sheet: Evidence" in selected
    assert "A1200=North" in selected
    assert "B1200=947" in selected
    assert "only an overview" not in selected
    assert hits["/report.xlsx"] == 1


class _DocxContinuationLlm:
    def __init__(self, url: str) -> None:
        self.url = url
        self.calls: list[dict[str, Any]] = []
        self.offset: int | None = None

    async def __call__(
        self,
        messages: list[dict[str, Any]],
        tools: list[dict[str, Any]] | None,
        *,
        tool_choice: str | None = None,
    ) -> Any:
        self.calls.append(
            {"messages": copy.deepcopy(messages), "tools": copy.deepcopy(tools), "tool_choice": tool_choice}
        )
        step = len(self.calls) - 1
        if step == 0:
            return _response(tool_calls=[_plan_call()])
        if step == 1:
            return _response(tool_calls=[_tool_call("page1", "fetch", {"url": self.url})])
        if step == 2:
            first = next(message["content"] for message in reversed(messages) if message.get("tool_call_id") == "page1")
            match = _CONTINUATION_RE.search(first)
            assert match is not None
            self.offset = int(match.group("offset"))
            return _response(tool_calls=[_tool_call("page2", "fetch", {"url": self.url, "start_char": self.offset})])
        return _response(tool_calls=[_tool_call("done", "conclude", {"gap_accounting": _accounting("g1")})])


@pytest.mark.asyncio
async def test_docx_order_and_dispatch_pagination_preserve_every_character_within_total_cap(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    document = Document()
    document.add_paragraph("BEFORE TABLE")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "OUTER CELL"
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "NESTED CELL"
    document.add_paragraph("AFTER TABLE")
    for index in range(180):
        document.add_paragraph(f"paragraph-{index:03d} " + "z" * 50)
    monkeypatch.setattr(guard, "is_public_http_url", lambda _url: _async_true())

    async with (
        _local_source_server(
            {
                "/report.docx": (
                    _docx_bytes(document),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            }
        ) as (urls, hits),
        aiohttp.ClientSession() as session,
    ):
        tools = agentic_tools.build_gap_fill_tools(
            "document ordering", ctx=agentic_tools.question_ladder_context(session=session)
        )
        llm = _DocxContinuationLlm(urls["/report.docx"])
        result = await run_agentic_loop(
            "system",
            "brief",
            tools,
            _config(max_steps=5, max_conclude_gate_rejections=0),
            llm_call=llm,
        )

    first = _tool_content(result, "page1")
    second = _tool_content(result, "page2")
    assert len(first) <= 8_000
    assert len(second) <= 8_000
    match = _CONTINUATION_RE.search(first)
    assert match is not None
    assert llm.offset == int(match.group("offset"))
    first_body = first.split("\n\n", 1)[1][: match.start() - first.index("\n\n") - 2]
    second_body = second.split("\n\n", 1)[1]
    parsed = run_cache.local_source_for(urls["/report.docx"])
    assert parsed is not None
    complete_text = source_text(select_source_sections(parsed))
    assert first_body + second_body == complete_text
    assert complete_text.index("BEFORE TABLE") < complete_text.index("OUTER CELL")
    assert complete_text.index("OUTER CELL") < complete_text.index("NESTED CELL")
    assert complete_text.index("NESTED CELL") < complete_text.index("AFTER TABLE")
    assert hits["/report.docx"] == 1


@pytest.mark.asyncio
async def test_resolution_policy_automatically_selects_late_relevant_passage_from_real_archive(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    rows = ["region,value", *(f"region-{index},{index}" for index in range(3000)), "North,947"]
    archive = _zip_bytes({"results.csv": ("\n".join(rows) + "\n").encode()})
    monkeypatch.setattr(guard, "is_public_http_url", lambda _url: _async_true())

    async with (
        _local_source_server({"/resolution.zip": (archive, "application/zip")}) as (urls, hits),
        aiohttp.ClientSession() as session,
    ):
        result = await fetch_url(
            urls["/resolution.zip"],
            policy=RESOLUTION_SOURCE_POLICY,
            ctx=LadderContext(
                query="North value",
                session=session,
                host_sems={},
            ),
        )

    assert result.status == "success"
    assert result.navigation_only is False
    assert "North" in result.text
    assert "947" in result.text
    assert "Member: results.csv" in result.text
    assert len(result.text) <= 6_000
    assert hits["/resolution.zip"] == 1


class _ImageEvidenceLlm:
    def __init__(self, page_url: str, image_url: str) -> None:
        self.page_url = page_url
        self.image_url = image_url
        self.calls: list[dict[str, Any]] = []

    async def __call__(
        self,
        messages: list[dict[str, Any]],
        tools: list[dict[str, Any]] | None,
        *,
        tool_choice: str | None = None,
    ) -> Any:
        self.calls.append(
            {"messages": copy.deepcopy(messages), "tools": copy.deepcopy(tools), "tool_choice": tool_choice}
        )
        step = len(self.calls) - 1
        if step == 0:
            return _response(tool_calls=[_plan_call()])
        if step == 1:
            return _response(tool_calls=[_tool_call("page", "fetch", {"url": self.page_url})])
        if step == 2:
            page_result = next(message["content"] for message in messages if message.get("tool_call_id") == "page")
            assert self.image_url in page_result
            return _response(
                tool_calls=[_tool_call("view", "view_image", {"url": self.image_url, "crop": [6, 0, 12, 8]})]
            )
        if step == 3:
            view_result = next(message["content"] for message in messages if message.get("tool_call_id") == "view")
            image_id_match = re.search(r"image_id=([0-9a-f]+)", view_result)
            assert image_id_match is not None
            return _response(
                tool_calls=[
                    _tool_call(
                        "record",
                        "record_findings",
                        {
                            "findings": [
                                {
                                    "claim": "The cropped chart region is yellow.",
                                    "source_url": self.image_url,
                                    "evidence_kind": "image",
                                    "image_id": image_id_match.group(1),
                                    "visual_observation": "The selected right half is yellow.",
                                    "retrieved_how": "view_image",
                                }
                            ]
                        },
                    )
                ]
            )
        return _response(tool_calls=[_tool_call("done", "conclude", {"gap_accounting": _accounting("g1")})])


def _image_blocks(call: dict[str, Any]) -> list[dict[str, Any]]:
    return [
        block
        for message in call["messages"]
        if isinstance(message.get("content"), list)
        for block in message["content"]
        if block.get("type") == "image_url"
    ]


class _DirectImageLlm:
    def __init__(self, image_urls: list[str]) -> None:
        self.image_urls = image_urls
        self.calls: list[dict[str, Any]] = []

    async def __call__(
        self,
        messages: list[dict[str, Any]],
        tools: list[dict[str, Any]] | None,
        *,
        tool_choice: str | None = None,
    ) -> Any:
        self.calls.append(
            {"messages": copy.deepcopy(messages), "tools": copy.deepcopy(tools), "tool_choice": tool_choice}
        )
        step = len(self.calls) - 1
        if step == 0:
            return _response(tool_calls=[_plan_call()])
        if step == 1:
            return _response(tool_calls=[_tool_call("fetch_direct", "fetch", {"url": self.image_urls[0]})])

        expected_views = min(step - 1, 4)
        assert len(_image_blocks(self.calls[-1])) == expected_views
        if step <= 5:
            assert self.image_urls[step - 2] in repr(messages)
        if step == 2:
            return _response(
                tool_calls=[
                    _tool_call(
                        "read_direct",
                        "read_document",
                        {"url": self.image_urls[1], "ask": "Read this chart"},
                    )
                ]
            )
        if step == 3:
            return _response(tool_calls=[_tool_call("view_explicit", "view_image", {"url": self.image_urls[2]})])
        if step == 4:
            return _response(tool_calls=[_tool_call("fetch_fourth", "fetch", {"url": self.image_urls[3]})])
        if step == 5:
            return _response(
                tool_calls=[
                    _tool_call(
                        "over_limit",
                        "read_document",
                        {"url": self.image_urls[4], "ask": "Read this fifth chart"},
                    )
                ]
            )
        return _response(tool_calls=[_tool_call("done", "conclude", {"gap_accounting": _accounting("g1")})])


@pytest.mark.asyncio
async def test_direct_image_tools_deliver_pixels_next_turn_and_share_four_view_budget(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(guard, "is_public_http_url", lambda _url: _async_true())
    paid_reader = MagicMock(side_effect=AssertionError("direct image reads must never reach Gemini"))
    monkeypatch.setattr(agentic_tools, "_run_document_read_sync", paid_reader)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    colors = [(255, 0, 0), (0, 255, 0), (0, 0, 255), (255, 255, 0), (255, 0, 255)]
    routes = {
        f"/chart-{index}.png": (_solid_png_bytes(color), "image/png") for index, color in enumerate(colors, start=1)
    }

    async with _local_source_server(routes) as (urls, hits), aiohttp.ClientSession() as session:
        image_urls = [urls[f"/chart-{index}.png"] for index in range(1, 6)]
        tools = agentic_tools.build_gap_fill_tools(
            "Compare the five charts", ctx=agentic_tools.question_ladder_context(session=session)
        )
        llm = _DirectImageLlm(image_urls)
        result = await run_agentic_loop(
            "system",
            "brief",
            tools,
            _config(wall_deadline_s=10.0, max_steps=8, max_conclude_gate_rejections=0),
            llm_call=llm,
        )

    for call_id in ("fetch_direct", "read_direct", "view_explicit", "fetch_fourth"):
        content = _tool_content(result, call_id)
        assert "status: ok" in content
        assert "method: image_local" in content
    refused = _tool_content(result, "over_limit")
    assert "status: error" in refused
    assert "method: image_local" in refused
    assert "already used its 4 distinct image views" in refused
    assert len(result.image_views) == 4
    assert [len(_image_blocks(call)) for call in llm.calls[2:7]] == [1, 2, 3, 4, 4]
    assert hits == Counter(dict.fromkeys(routes, 1))
    paid_reader.assert_not_called()


@pytest.mark.asyncio
async def test_html_image_lead_is_cropped_and_delivered_to_same_driver_before_visual_finding(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(guard, "is_public_http_url", lambda _url: _async_true())
    image_body = _png_bytes()

    routes = {
        "/index.html": (b"placeholder", "text/html"),
        "/chart.png": (image_body, "image/png"),
    }
    async with _local_source_server(routes) as (urls, hits), aiohttp.ClientSession() as session:
        routes["/index.html"] = (
            (
                "<html><body><main><h1>Hospitalizations report</h1>"
                "<p>The official chart below contains the final monthly observations and source notes.</p>"
                f'<figure><img src="{urls["/chart.png"]}" alt="Monthly hospitalizations through September">'
                "<figcaption>Official monthly hospitalizations</figcaption></figure>"
                "</main></body></html>"
            ).encode(),
            "text/html",
        )
        tools = agentic_tools.build_gap_fill_tools(
            "September hospitalizations", ctx=agentic_tools.question_ladder_context(session=session)
        )
        llm = _ImageEvidenceLlm(urls["/index.html"], urls["/chart.png"])
        result = await run_agentic_loop(
            "system",
            "brief",
            tools,
            _config(max_steps=6, max_conclude_gate_rejections=0),
            llm_call=llm,
        )

    assert hits == Counter({"/index.html": 1, "/chart.png": 1})
    assert len(result.image_views) == 1
    assert result.image_views[0].crop == (6, 0, 12, 8)
    assert (result.image_views[0].width, result.image_views[0].height) == (6, 8)
    assert len(_image_blocks(llm.calls[3])) == 1
    assert llm.calls[3]["tools"] == llm.calls[2]["tools"]
    assert len(llm.calls) == 5
    assert result.telemetry.findings_count == 1
    assert "The cropped chart region is yellow." in result.findings_markdown
    assert "data:image" not in repr(result.transcript)


@pytest.mark.asyncio
async def test_real_litellm_openrouter_boundary_accepts_materialized_pixels_on_configured_model(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    serialized_requests: list[tuple[str, dict[str, Any]]] = []

    async def openrouter_response(
        _handler: AsyncHTTPHandler,
        url: str,
        data: dict[str, Any] | str | bytes | None = None,
        **_kwargs: Any,
    ) -> httpx.Response:
        assert isinstance(data, str | bytes)
        body = json.loads(data)
        serialized_requests.append((url, body))
        request = httpx.Request("POST", url, content=data)
        return httpx.Response(
            200,
            json={
                "id": "chatcmpl-offline",
                "object": "chat.completion",
                "created": 1_700_000_000,
                "model": GAP_FILL_V2_DRIVER_MODEL,
                "choices": [
                    {
                        "index": 0,
                        "finish_reason": "stop",
                        "message": {"role": "assistant", "content": "ok"},
                    }
                ],
                "usage": {"prompt_tokens": 1, "completion_tokens": 1, "total_tokens": 2},
            },
            request=request,
        )

    monkeypatch.setattr(agentic_llm, "should_route_via_donated_key", lambda _model: False)
    monkeypatch.setattr(AsyncHTTPHandler, "post", openrouter_response)
    monkeypatch.setattr(litellm, "drop_params", True)
    monkeypatch.setenv("OPENROUTER_API_KEY", "sk-offline-wire-format-test")
    image = ImageView(
        image_id="pixel-sha",
        source_url="https://charts.example/chart.png",
        final_url="https://charts.example/chart.png",
        source_sha256="source-sha",
        original_width=12,
        original_height=8,
        width=6,
        height=8,
        crop=(6, 0, 12, 8),
        png_bytes=_png_bytes(),
    )
    tools_json = [
        {
            "type": "function",
            "function": {
                "name": "conclude",
                "description": "finish",
                "parameters": {"type": "object", "properties": {}},
            },
        }
    ]
    references = [
        {
            "role": "assistant",
            "content": None,
            "tool_calls": [
                {
                    "id": "view1",
                    "type": "function",
                    "function": {"name": "view_image", "arguments": '{"url":"https://charts.example/chart.png"}'},
                }
            ],
        },
        {"role": "tool", "tool_call_id": "view1", "content": "image_id=pixel-sha"},
        image_reference_message([image.image_id]),
    ]
    call = bind_image_materialization(
        agentic_llm.build_default_llm_call(_config(model=GAP_FILL_V2_DRIVER_MODEL)),
        {image.image_id: image},
        {image.image_id: {image.source_url}},
    )
    response = await call(references, tools_json)

    assert response.choices[0].message.content == "ok"
    assert len(serialized_requests) == 1
    url, body = serialized_requests[0]
    assert url == "https://openrouter.ai/api/v1/chat/completions"
    assert body["model"] == GAP_FILL_V2_DRIVER_MODEL
    assert body["tools"] == tools_json
    assert [message["role"] for message in body["messages"]] == ["assistant", "tool", "user"]
    assert body["messages"][0]["tool_calls"][0]["function"]["name"] == "view_image"
    assert body["messages"][1]["tool_call_id"] == "view1"
    image_content = body["messages"][2]["content"]
    assert [block["type"] for block in image_content] == ["text", "text", "image_url"]
    assert image_content[-1]["image_url"]["url"].startswith("data:image/png;base64,")
    assert references[-1]["content"][-1] == {"type": "image_reference", "image_id": image.image_id}


async def _async_true() -> bool:
    return True
