"""Bounded, local-only readers for office documents and ZIP archives.

The caller owns archive inventory policy, member selection, and digest rendering. This module
only turns already-held bytes into complete labeled sections, with limits applied before a
partial result can escape.
"""

from __future__ import annotations

import csv
import re
import struct
import zlib
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO, StringIO
from typing import Any, Literal
from zipfile import BadZipFile, LargeZipFile, ZipFile, ZipInfo

import xlrd
from docx import Document
from docx.drawing import Drawing
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.pagebreak import RenderedPageBreak
from docx.text.paragraph import Paragraph
from lxml.etree import XMLSyntaxError
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.exceptions import InvalidFileException
from xlrd.biffh import XLRDError

from metaculus_bot.constants import (
    LOCAL_SOURCE_CSV_SNIFF_CHARS,
    LOCAL_SOURCE_MAX_CELLS,
    LOCAL_SOURCE_MAX_CHARS,
    LOCAL_SOURCE_MAX_ENTRIES,
    LOCAL_SOURCE_MAX_EXPANDED_BYTES,
    LOCAL_SOURCE_MAX_SHEETS,
    LOCAL_SOURCE_READ_CHUNK_BYTES,
)

_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ZIP_MAGICS = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_XLSM_MIME = "application/vnd.ms-excel.sheet.macroenabled.12"
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLS_MIME = "application/vnd.ms-excel"
_ZIP_MIMES = {"application/zip", "application/x-zip-compressed", "multipart/x-zip"}
_TEXT_SUFFIXES = {".txt", ".md", ".json", ".xml", ".log"}
_CSV_SUFFIXES = {".csv", ".tsv"}
_XLSX_SUFFIXES = {".xlsx", ".xlsm"}
_DOCX_SUFFIXES = {".docx"}
_XLS_SUFFIXES = {".xls"}
_NESTED_ARCHIVE_SUFFIXES = {".zip", ".7z", ".rar", ".tar", ".gz", ".bz2", ".xz"}
_SHEET_CACHE_NOTICE = "Formula results use saved workbook values; formulas are not evaluated locally."
_WORD_OMISSION_NOTICE = "Visual content omitted. Only body paragraphs, tables, headers, and footers are read; other unsupported Word parts and embedded objects are omitted."


@dataclass(frozen=True, slots=True)
class SourceSection:
    """One complete, location-labeled text section from a source."""

    member: str | None
    sheet: str | None
    text: str
    rows: int | None = None
    columns: int | None = None


@dataclass(frozen=True, slots=True)
class SourceMember:
    """One archive entry and whether this parser could read it."""

    name: str
    size_bytes: int
    readable: bool
    kind: str | None = None
    reason: str | None = None


@dataclass(frozen=True, slots=True)
class ParsedSource:
    """The extracted sections and visible archive-member inventory."""

    kind: Literal["archive", "workbook", "word"]
    sections: tuple[SourceSection, ...]
    members: tuple[SourceMember, ...] = ()
    notices: tuple[str, ...] = ()

    @property
    def retained_bytes(self) -> int:
        """Estimate UTF-8 text plus the metadata retained beside it."""
        total = sum(len(section.text.encode("utf-8")) for section in self.sections)
        for section in self.sections:
            total += _metadata_size(section.member, section.sheet)
        for member in self.members:
            total += _metadata_size(
                member.name,
                str(member.size_bytes),
                str(member.readable),
                member.kind,
                member.reason,
            )
        return total + sum(len(notice.encode("utf-8")) for notice in self.notices)


class SourceReadError(ValueError):
    """A malformed, encrypted, unsupported, or over-limit source that cannot be read."""

    def __init__(self, reason: str) -> None:
        super().__init__(reason)
        self.reason = reason


@dataclass(slots=True)
class _ReadBudget:
    expanded_bytes: int = 0
    archive_entries: int = 0
    cells: int = 0
    sheets: int = 0
    characters: int = 0

    def add_expanded_bytes(self, size: int) -> None:
        self.expanded_bytes += size
        if self.expanded_bytes > LOCAL_SOURCE_MAX_EXPANDED_BYTES:
            raise SourceReadError("expanded byte limit exceeded")

    def add_entries(self, count: int) -> None:
        self.archive_entries += count
        if self.archive_entries > LOCAL_SOURCE_MAX_ENTRIES:
            raise SourceReadError("archive entry limit exceeded")

    def add_cells(self, count: int) -> None:
        self.cells += count
        if self.cells > LOCAL_SOURCE_MAX_CELLS:
            raise SourceReadError("workbook or table cell limit exceeded")

    def add_sheets(self, count: int) -> None:
        self.sheets += count
        if self.sheets > LOCAL_SOURCE_MAX_SHEETS:
            raise SourceReadError("workbook sheet limit exceeded")

    def add_line(self, lines: list[str], line: str) -> None:
        new_size = len(line) + (1 if lines else 0)
        self.add_characters(new_size)
        lines.append(line)

    def add_characters(self, count: int) -> None:
        self.characters += count
        if self.characters > LOCAL_SOURCE_MAX_CHARS:
            raise SourceReadError("extracted character limit exceeded")


def _metadata_size(*values: str | None) -> int:
    return sum(len(value.encode("utf-8")) for value in values if value is not None) + 2 * len(values)


def _normalized_mime(content_type: str) -> str:
    return content_type.partition(";")[0].strip().lower()


def is_local_source(body: bytes, content_type: str) -> bool:
    """Recognize a supported local container from a short signature or MIME value."""
    mime = _normalized_mime(content_type)
    if mime in _ZIP_MIMES or mime in {_XLSX_MIME, _XLSM_MIME, _DOCX_MIME, _XLS_MIME}:
        return True
    if mime == "application/msword":
        return True
    return body.startswith(_OLE_MAGIC) or _has_zip_signature(body)


def parse_source(body: bytes, content_type: str) -> ParsedSource:
    """Parse supported local bytes without network access or filesystem extraction."""
    mime = _normalized_mime(content_type)
    budget = _ReadBudget()

    if body.startswith(_OLE_MAGIC):
        return _parse_ole_source(body, mime=mime, budget=budget)
    zip_names = _zip_names(body)
    if zip_names is not None:
        return _parse_zip_source(body, mime=mime, names=zip_names, budget=budget)
    return _parse_source_by_mime(body, mime=mime, budget=budget)


def _parse_ole_source(body: bytes, *, mime: str, budget: _ReadBudget) -> ParsedSource:
    if mime in {_DOCX_MIME, "application/msword"}:
        raise SourceReadError("encrypted or legacy Word containers are unsupported")
    if mime in {_XLSX_MIME, _XLSM_MIME}:
        raise SourceReadError("encrypted workbook containers are unsupported")
    return _parse_xls(body, member=None, budget=budget)


def _parse_zip_source(body: bytes, *, mime: str, names: set[str], budget: _ReadBudget) -> ParsedSource:
    has_content_types = "[Content_Types].xml" in names
    is_xlsx = has_content_types and "xl/workbook.xml" in names
    is_docx = has_content_types and "word/document.xml" in names
    if is_xlsx and is_docx:
        raise SourceReadError("malformed office package contains multiple document roots")
    if is_xlsx:
        return _parse_xlsx(body, member=None, budget=budget)
    if is_docx:
        return _parse_word(body, member=None, budget=budget)
    if mime in {_XLSX_MIME, _XLSM_MIME}:
        return _parse_xlsx(body, member=None, budget=budget)
    if mime == _DOCX_MIME:
        return _parse_word(body, member=None, budget=budget)
    if mime == _XLS_MIME:
        raise SourceReadError("malformed legacy workbook: expected an OLE workbook")
    return _parse_archive(body, budget=budget)


def _parse_source_by_mime(body: bytes, *, mime: str, budget: _ReadBudget) -> ParsedSource:
    if mime in {_XLSX_MIME, _XLSM_MIME}:
        return _parse_xlsx(body, member=None, budget=budget)
    if mime == _DOCX_MIME:
        return _parse_word(body, member=None, budget=budget)
    if mime == _XLS_MIME:
        raise SourceReadError("malformed legacy workbook: expected an OLE workbook")
    if mime in _ZIP_MIMES or _has_zip_signature(body):
        return _parse_archive(body, budget=budget)
    if mime == "application/msword":
        raise SourceReadError("legacy Word binary documents are unsupported")
    raise SourceReadError("unsupported source type")


def _zip_names(body: bytes) -> set[str] | None:
    if not _has_zip_signature(body):
        return None
    try:
        with ZipFile(BytesIO(body)) as archive:
            infos = archive.infolist()
            if len(infos) > LOCAL_SOURCE_MAX_ENTRIES:
                raise SourceReadError("archive entry limit exceeded")
            return {info.filename for info in infos}
    except SourceReadError:
        raise
    except (BadZipFile, LargeZipFile, OSError, EOFError):
        return None


def _has_zip_signature(body: bytes) -> bool:
    return body.startswith(_ZIP_MAGICS)


def _register_zip(archive: ZipFile, budget: _ReadBudget, *, label: str) -> list[ZipInfo]:
    infos = archive.infolist()
    budget.add_entries(len(infos))
    declared_size = sum(info.file_size for info in infos if not info.is_dir())
    if budget.expanded_bytes + declared_size > LOCAL_SOURCE_MAX_EXPANDED_BYTES:
        raise SourceReadError("expanded byte limit exceeded by archive metadata")
    for info in infos:
        if info.flag_bits & 1:
            raise SourceReadError(f"encrypted {label} entry is unsupported: {info.filename}")
    return infos


def _read_zip_member(
    archive: ZipFile,
    info: ZipInfo,
    budget: _ReadBudget,
    *,
    capture: bool,
    label: str,
) -> bytes:
    if info.flag_bits & 1:
        raise SourceReadError(f"encrypted {label} entry is unsupported: {info.filename}")
    chunks: list[bytes] = []
    try:
        with archive.open(info) as member_stream:
            while chunk := member_stream.read(LOCAL_SOURCE_READ_CHUNK_BYTES):
                budget.add_expanded_bytes(len(chunk))
                if capture:
                    chunks.append(chunk)
    except SourceReadError:
        raise
    except (BadZipFile, RuntimeError, NotImplementedError, OSError, EOFError, zlib.error) as error:
        raise SourceReadError(f"malformed {label} entry {info.filename}: {error}") from error
    return b"".join(chunks)


def _validate_office_package(body: bytes, budget: _ReadBudget, *, expected_root: str) -> None:
    try:
        with ZipFile(BytesIO(body)) as archive:
            infos = _register_zip(archive, budget, label="office package")
            names = {info.filename for info in infos}
            if "[Content_Types].xml" not in names or expected_root not in names:
                raise SourceReadError(f"malformed office package: missing {expected_root}")
            for info in infos:
                if not info.is_dir():
                    _read_zip_member(archive, info, budget, capture=False, label="office package")
    except SourceReadError:
        raise
    except (BadZipFile, LargeZipFile, OSError, EOFError) as error:
        raise SourceReadError(f"malformed office package: {error}") from error


def _parse_archive(body: bytes, *, budget: _ReadBudget) -> ParsedSource:
    sections: list[SourceSection] = []
    members: list[SourceMember] = []
    notices: list[str] = []
    try:
        with ZipFile(BytesIO(body)) as archive:
            infos = _register_zip(archive, budget, label="archive")
            _parse_archive_entries(
                archive,
                infos,
                budget,
                sections=sections,
                members=members,
                notices=notices,
            )
    except SourceReadError:
        raise
    except (BadZipFile, LargeZipFile, OSError, EOFError) as error:
        raise SourceReadError(f"malformed archive: {error}") from error

    return ParsedSource("archive", tuple(sections), tuple(members), tuple(notices))


def _parse_archive_entries(
    archive: ZipFile,
    infos: list[ZipInfo],
    budget: _ReadBudget,
    *,
    sections: list[SourceSection],
    members: list[SourceMember],
    notices: list[str],
) -> None:
    for info in infos:
        member_sections, source_member, member_notices = _parse_archive_entry(archive, info, budget)
        sections.extend(member_sections)
        members.append(source_member)
        notices.extend(notice for notice in member_notices if notice not in notices)


def _parse_archive_entry(
    archive: ZipFile,
    info: ZipInfo,
    budget: _ReadBudget,
) -> tuple[list[SourceSection], SourceMember, tuple[str, ...]]:
    name = info.filename
    if info.is_dir():
        return (
            [],
            SourceMember(
                name,
                info.file_size,
                readable=False,
                kind="directory",
                reason="directory has no file content",
            ),
            (),
        )
    suffix = _suffix(name)
    if suffix in _NESTED_ARCHIVE_SUFFIXES:
        _read_zip_member(archive, info, budget, capture=False, label="archive")
        member = SourceMember(
            name,
            info.file_size,
            readable=False,
            kind="archive",
            reason="nested archives are not recursively parsed",
        )
        return [], member, ()

    is_candidate = suffix in (_TEXT_SUFFIXES | _CSV_SUFFIXES | _XLSX_SUFFIXES | _DOCX_SUFFIXES | _XLS_SUFFIXES)
    data = _read_zip_member(archive, info, budget, capture=is_candidate, label="archive")
    if not is_candidate:
        member = SourceMember(name, info.file_size, readable=False, reason="unsupported archive member type")
        return [], member, ()
    try:
        sections, member_kind, notices = _parse_archive_member(data, name=name, suffix=suffix, budget=budget)
    except SourceReadError as error:
        if _is_limit_error(error):
            raise
        member = SourceMember(
            name,
            info.file_size,
            readable=False,
            kind=_kind_for_suffix(suffix),
            reason=error.reason,
        )
        return [], member, ()
    member = SourceMember(name, info.file_size, readable=True, kind=member_kind)
    return sections, member, notices


def _parse_archive_member(
    body: bytes,
    *,
    name: str,
    suffix: str,
    budget: _ReadBudget,
) -> tuple[list[SourceSection], str, tuple[str, ...]]:
    if suffix in _XLSX_SUFFIXES:
        parsed = _parse_xlsx(body, member=name, budget=budget)
        return list(parsed.sections), "workbook", parsed.notices
    if suffix in _XLS_SUFFIXES:
        parsed = _parse_xls(body, member=name, budget=budget)
        return list(parsed.sections), "workbook", parsed.notices
    if suffix in _DOCX_SUFFIXES:
        parsed = _parse_word(body, member=name, budget=budget)
        return list(parsed.sections), "word", parsed.notices
    if suffix in _CSV_SUFFIXES:
        section = _parse_delimited(body, name=name, suffix=suffix, budget=budget)
        return [section], "csv" if suffix == ".csv" else "tsv", ()
    if suffix in _TEXT_SUFFIXES:
        section = _parse_text_member(body, name=name, budget=budget)
        return [section], "text", ()
    raise SourceReadError("unsupported archive member type")


def _parse_delimited(body: bytes, *, name: str, suffix: str, budget: _ReadBudget) -> SourceSection:
    try:
        text = body.decode("utf-8-sig")
        delimiter = _csv_delimiter(text, suffix)
        records = csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True)
    except UnicodeDecodeError as error:
        raise SourceReadError(f"malformed {suffix[1:]} member: {error}") from error

    lines: list[str] = []
    max_columns = 0
    row_count = 0
    try:
        for row_count, record in enumerate(records, start=1):
            budget.add_cells(len(record))
            max_columns = max(max_columns, len(record))
            values = "\t".join(_safe_cell_text(value) if value else "[blank]" for value in record)
            location = f"{_cell_name(row_count, 1)}:{_cell_name(row_count, len(record))}" if record else "empty"
            budget.add_line(lines, f"row {row_count} ({location}): {values}")
    except csv.Error as error:
        raise SourceReadError(f"malformed {suffix[1:]} member: {error}") from error
    column_locations = " ".join(_column_name(column) for column in range(1, max_columns + 1))
    column_line = f"columns: {column_locations}"
    budget.add_characters(len(column_line) + (1 if lines else 0))
    lines.insert(0, column_line)
    return SourceSection(name, None, "\n".join(lines), rows=row_count, columns=max_columns)


def _csv_delimiter(text: str, suffix: str) -> str:
    if suffix == ".tsv":
        return "\t"
    sample_lines: list[str] = []
    sample_chars = 0
    for source_line in StringIO(text):
        line = source_line.rstrip("\r\n")
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        line_size = len(line) + 1
        if sample_chars + line_size > LOCAL_SOURCE_CSV_SNIFF_CHARS:
            if not sample_lines:
                sample_lines.append(line[:LOCAL_SOURCE_CSV_SNIFF_CHARS])
            break
        sample_lines.append(line)
        sample_chars += line_size
    sample = "\n".join(sample_lines)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        sniffed_delimiter = ","
    else:
        sniffed_delimiter = dialect.delimiter
    return max(",;\t", key=lambda delimiter: _delimiter_score(sample, delimiter, sniffed_delimiter))


def _delimiter_score(sample: str, delimiter: str, sniffed_delimiter: str) -> tuple[int, int, int, int]:
    try:
        rows = list(csv.reader(StringIO(sample, newline=""), delimiter=delimiter, strict=True))
    except csv.Error:
        return 0, 0, 0, int(delimiter == sniffed_delimiter)
    widths = [len(row) for row in rows if row]
    if not widths:
        return 0, 0, 0, int(delimiter == sniffed_delimiter)
    consistent = int(min(widths) == max(widths))
    return min(widths), consistent, sum(widths), int(delimiter == sniffed_delimiter)


def _parse_text_member(body: bytes, *, name: str, budget: _ReadBudget) -> SourceSection:
    try:
        text = body.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise SourceReadError(f"text member is not UTF-8: {error}") from error
    lines: list[str] = []
    for line_number, line in enumerate(StringIO(text), start=1):
        budget.add_line(lines, f"line {line_number}: {_safe_cell_text(line)}")
    return SourceSection(name, None, "\n".join(lines), rows=len(lines), columns=None)


def _parse_xlsx(body: bytes, *, member: str | None, budget: _ReadBudget) -> ParsedSource:
    try:
        _validate_office_package(body, budget, expected_root="xl/workbook.xml")
    except SourceReadError as error:
        if _is_limit_error(error) or "encrypted" in error.reason:
            raise
        raise SourceReadError(f"malformed workbook: {error.reason}") from error
    values_workbook = _load_xlsx_workbook(body, data_only=True)
    try:
        formula_workbook = _load_xlsx_workbook(body, data_only=False)
    except SourceReadError:
        values_workbook.close()
        raise

    try:
        if len(values_workbook.worksheets) != len(formula_workbook.worksheets):
            raise SourceReadError("malformed workbook: inconsistent sheet metadata")
        budget.add_sheets(len(values_workbook.worksheets))
        parsed_sheets = [
            _parse_xlsx_sheet(value_sheet, formula_sheet, member=member, budget=budget)
            for value_sheet, formula_sheet in zip(
                values_workbook.worksheets,
                formula_workbook.worksheets,
                strict=True,
            )
        ]
        sections = [section for section, _has_formula in parsed_sheets]
        has_formula = any(sheet_has_formula for _section, sheet_has_formula in parsed_sheets)
    except SourceReadError:
        raise
    except (
        BadZipFile,
        KeyError,
        ValueError,
        TypeError,
        IndexError,
        OSError,
        EOFError,
        XMLSyntaxError,
        zlib.error,
    ) as error:
        raise SourceReadError(f"malformed workbook: {error}") from error
    finally:
        values_workbook.close()
        formula_workbook.close()
    notices = (_SHEET_CACHE_NOTICE,) if has_formula else ()
    return ParsedSource("workbook", tuple(sections), notices=notices)


def _load_xlsx_workbook(body: bytes, *, data_only: bool) -> Any:
    try:
        return load_workbook(BytesIO(body), read_only=True, data_only=data_only, keep_links=False)
    except (
        InvalidFileException,
        BadZipFile,
        KeyError,
        ValueError,
        OSError,
        EOFError,
        XMLSyntaxError,
        zlib.error,
    ) as error:
        raise SourceReadError(f"malformed workbook: {error}") from error


def _parse_xlsx_sheet(
    values_sheet: Any,
    formula_sheet: Any,
    *,
    member: str | None,
    budget: _ReadBudget,
) -> tuple[SourceSection, bool]:
    min_column, min_row, max_column, max_row = _sheet_boundaries(values_sheet, formula_sheet)
    rows = max_row - min_row + 1
    columns = max_column - min_column + 1
    budget.add_cells(rows * columns)
    lines: list[str] = []
    budget.add_line(lines, f"Sheet: {values_sheet.title}")
    value_rows = values_sheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_column,
        max_col=max_column,
    )
    formula_rows = formula_sheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_column,
        max_col=max_column,
    )
    has_formula = False
    for row_number, (value_row, formula_row) in enumerate(zip(value_rows, formula_rows, strict=True), start=min_row):
        row_text, row_has_formula = _render_xlsx_row(row_number, value_row, formula_row, min_column=min_column)
        has_formula = has_formula or row_has_formula
        budget.add_line(lines, row_text)
    return SourceSection(member, values_sheet.title, "\n".join(lines), rows=rows, columns=columns), has_formula


def _sheet_boundaries(values_sheet: Any, formula_sheet: Any) -> tuple[int, int, int, int]:
    try:
        values_dimension = values_sheet.calculate_dimension(force=True)
        formula_dimension = formula_sheet.calculate_dimension(force=True)
        if values_dimension != formula_dimension:
            raise SourceReadError(f"malformed workbook: inconsistent dimensions on {values_sheet.title}")
        min_column, min_row, max_column, max_row = range_boundaries(values_dimension)
        if min_column is None or min_row is None or max_column is None or max_row is None:
            raise SourceReadError(f"malformed workbook: empty dimensions on {values_sheet.title}")
        return min_column, min_row, max_column, max_row
    except SourceReadError:
        raise
    except (ValueError, TypeError, KeyError, XMLSyntaxError, OSError, zlib.error) as error:
        raise SourceReadError(f"malformed workbook sheet {values_sheet.title}: {error}") from error


def _render_xlsx_row(
    row_number: int,
    value_row: tuple[Any, ...],
    formula_row: tuple[Any, ...],
    *,
    min_column: int,
) -> tuple[str, bool]:
    cells = []
    has_formula = False
    for column_number, (value_cell, formula_cell) in enumerate(
        zip(value_row, formula_row, strict=True), start=min_column
    ):
        rendered_value, is_formula = _format_xlsx_cell(value_cell, formula_cell)
        has_formula = has_formula or is_formula
        cells.append(f"{_cell_name(row_number, column_number)}={rendered_value}")
    return f"row {row_number}: {' | '.join(cells)}", has_formula


def _format_xlsx_cell(value_cell: Any, formula_cell: Any) -> tuple[str, bool]:
    if formula_cell.data_type != "f":
        return _format_openpyxl_value(value_cell), False
    if value_cell.value is None:
        return f"{_safe_cell_text(str(formula_cell.value))} [cached value unavailable]", True
    return f"{_format_openpyxl_value(value_cell)} [cached result]", True


def _format_openpyxl_value(cell: object) -> str:
    value = cell.value  # type: ignore[attr-defined]
    if value is None:
        return "[blank]"
    if cell.data_type == "e":  # type: ignore[attr-defined]
        return f"[Excel error: {_safe_cell_text(str(value))}]"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return _format_number(value, cell.number_format)  # type: ignore[attr-defined]
    return _safe_cell_text(str(value))


def _parse_xls(body: bytes, *, member: str | None, budget: _ReadBudget) -> ParsedSource:
    if not body.startswith(_OLE_MAGIC):
        raise SourceReadError("malformed legacy workbook: invalid OLE signature")
    try:
        workbook = xlrd.open_workbook(file_contents=body, on_demand=True, formatting_info=True)
    except XLRDError as error:
        if "encrypt" in str(error).lower() or "password" in str(error).lower():
            raise SourceReadError("encrypted legacy workbook is unsupported") from error
        raise SourceReadError(f"malformed legacy workbook: {error}") from error
    except (struct.error, ValueError, IndexError, KeyError, OSError, zlib.error) as error:
        raise SourceReadError(f"malformed legacy workbook: {error}") from error

    sections: list[SourceSection] = []
    try:
        budget.add_sheets(workbook.nsheets)
        for sheet_index in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(sheet_index)
            rows = sheet.nrows
            columns = sheet.ncols
            budget.add_cells(rows * columns)
            lines: list[str] = []
            budget.add_line(lines, f"Sheet: {sheet.name}")
            budget.add_line(lines, _SHEET_CACHE_NOTICE)
            for row_number in range(rows):
                rendered_cells: list[str] = []
                for column_number in range(columns):
                    cell = sheet.cell(row_number, column_number)
                    coordinate = _cell_name(row_number + 1, column_number + 1)
                    value = _format_xlrd_value(workbook, cell)
                    rendered_cells.append(f"{coordinate}={value}")
                budget.add_line(lines, f"row {row_number + 1}: {' | '.join(rendered_cells)}")
            sections.append(SourceSection(member, sheet.name, "\n".join(lines), rows=rows, columns=columns))
    finally:
        workbook.release_resources()
    return ParsedSource("workbook", tuple(sections), notices=(_SHEET_CACHE_NOTICE,))


def _format_xlrd_value(workbook: xlrd.book.Book, cell: xlrd.sheet.Cell) -> str:
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
        return "[blank]"
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return f"[Excel error: {xlrd.error_text_from_code.get(int(cell.value), cell.value)}]"
    if cell.ctype == xlrd.XL_CELL_DATE:
        if not isinstance(cell.value, (int, float)):
            raise SourceReadError("malformed workbook date value: expected a number")
        try:
            return xlrd.xldate_as_datetime(float(cell.value), workbook.datemode).isoformat()
        except (ValueError, OverflowError) as error:
            raise SourceReadError(f"malformed workbook date value: {error}") from error
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        if not isinstance(cell.value, (int, float)):
            raise SourceReadError("malformed workbook number value: expected a number")
        number_format = ""
        if cell.xf_index is not None:
            number_format_key = workbook.xf_list[cell.xf_index].format_key
            number_format = workbook.format_map[number_format_key].format_str
        return _format_number(float(cell.value), number_format)
    return _safe_cell_text(str(cell.value))


def _format_number(value: int | float, number_format: str) -> str:
    if _has_active_percent_token(number_format):
        first_format_section = number_format.split(";", maxsplit=1)[0]
        decimal_match = re.search(r"\.([0#?]+)[^%]*%", first_format_section)
        decimal_places = len(decimal_match.group(1)) if decimal_match else 0
        return f"{value:.{decimal_places}%}"
    if "%" in number_format:
        return f"{format(value, '.15g')} [literal percent format: {_safe_cell_text(number_format)}]"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return format(value, ".15g")


def _has_active_percent_token(number_format: str) -> bool:
    in_quoted_literal = False
    escaped = False
    for character in number_format:
        if escaped:
            escaped = False
        elif character == "\\":
            escaped = True
        elif character == '"':
            in_quoted_literal = not in_quoted_literal
        elif character == "%" and not in_quoted_literal:
            return True
    return False


def _parse_word(body: bytes, *, member: str | None, budget: _ReadBudget) -> ParsedSource:
    try:
        _validate_office_package(body, budget, expected_root="word/document.xml")
    except SourceReadError as error:
        if _is_limit_error(error) or "encrypted" in error.reason:
            raise
        raise SourceReadError(f"malformed Word document: {error.reason}") from error
    try:
        document = Document(BytesIO(body))
    except (
        PackageNotFoundError,
        BadZipFile,
        KeyError,
        ValueError,
        OSError,
        EOFError,
        XMLSyntaxError,
        zlib.error,
    ) as error:
        raise SourceReadError(f"malformed Word document: {error}") from error

    body_lines = _render_word_container(document.iter_inner_content(), budget)
    sections = [SourceSection(member, "body", "\n".join(body_lines))]
    seen_story_parts: set[str] = set()
    for section_number, document_section in enumerate(document.sections, start=1):
        story_variants = (
            ("header", document_section.header),
            ("first-page header", document_section.first_page_header),
            ("even-page header", document_section.even_page_header),
            ("footer", document_section.footer),
            ("first-page footer", document_section.first_page_footer),
            ("even-page footer", document_section.even_page_footer),
        )
        for location, story in story_variants:
            story_part_name = str(story.part.partname)
            if story_part_name in seen_story_parts:
                continue
            seen_story_parts.add(story_part_name)
            lines = _render_word_container(story.iter_inner_content(), budget)
            if lines:
                label = f"{location} section {section_number}"
                labeled_lines = [f"[{label}]"]
                budget.add_characters(len(labeled_lines[0]) + 1)
                labeled_lines.extend(lines)
                sections.append(SourceSection(member, label, "\n".join(labeled_lines)))
    return ParsedSource("word", tuple(sections), notices=(_WORD_OMISSION_NOTICE,))


def _render_word_container(content: Iterable[Paragraph | Table], budget: _ReadBudget) -> list[str]:
    lines: list[str] = []
    paragraph_number = 0
    table_number = 0
    for item in content:  # type: ignore[union-attr]
        if isinstance(item, Paragraph):
            paragraph_number += 1
            text, omissions = _paragraph_text(item)
            if text.strip() or omissions:
                rendered = text if text.strip() else "[blank]"
                rendered += "".join(f" [{omission}]" for omission in omissions)
                budget.add_line(lines, f"paragraph {paragraph_number}: {rendered}")
        elif isinstance(item, Table):
            table_number += 1
            lines.extend(_render_word_table(item, f"table {table_number}", budget))
    return lines


def _render_word_table(table: Table, label: str, budget: _ReadBudget) -> list[str]:
    lines: list[str] = []
    for row_number, row in enumerate(table.rows, start=1):  # type: ignore[attr-defined]
        budget.add_cells(len(row.cells))
        for column_number, cell in enumerate(row.cells, start=1):
            cell_label = f"{label} row {row_number} cell {column_number}"
            lines.extend(_render_word_cell(cell, cell_label, budget))
    return lines


def _render_word_cell(cell: Any, label: str, budget: _ReadBudget) -> list[str]:
    lines: list[str] = []
    for inner_item in cell.iter_inner_content():
        if isinstance(inner_item, Paragraph):
            text, omissions = _paragraph_text(inner_item)
            if text.strip() or omissions:
                rendered = text if text.strip() else "[blank]"
                rendered += "".join(f" [{omission}]" for omission in omissions)
                budget.add_line(lines, f"{label}: {rendered}")
        elif isinstance(inner_item, Table):
            lines.extend(_render_word_table(inner_item, f"{label} nested table", budget))
    return lines


def _paragraph_text(paragraph: Paragraph) -> tuple[str, tuple[str, ...]]:
    text = paragraph.text
    omissions: list[str] = []
    for run in paragraph.runs:
        for content in run.iter_inner_content():
            if isinstance(content, Drawing):
                omissions.append("visual content omitted")
            elif isinstance(content, RenderedPageBreak):
                omissions.append("rendered page break omitted")
    return text, tuple(omissions)


def _cell_name(row_number: int, column_number: int) -> str:
    return f"{_column_name(column_number)}{row_number}"


def _column_name(column_number: int) -> str:
    column_label = ""
    while column_number:
        column_number, remainder = divmod(column_number - 1, 26)
        column_label = chr(65 + remainder) + column_label
    return column_label


def _safe_cell_text(value: str) -> str:
    return value.replace("\r", "\\r").replace("\n", "\\n").replace("\t", "\\t")


def _suffix(name: str) -> str:
    basename = name.rsplit("/", maxsplit=1)[-1]
    if "." not in basename:
        return ""
    return f".{basename.rsplit('.', maxsplit=1)[-1].lower()}"


def _kind_for_suffix(suffix: str) -> str | None:
    if suffix in _XLSX_SUFFIXES | _XLS_SUFFIXES:
        return "workbook"
    if suffix in _DOCX_SUFFIXES:
        return "word"
    if suffix == ".csv":
        return "csv"
    if suffix == ".tsv":
        return "tsv"
    if suffix in _TEXT_SUFFIXES:
        return "text"
    return None


def _is_limit_error(error: SourceReadError) -> bool:
    return any(term in error.reason for term in ("limit exceeded", "limit", "too many"))
